"""
Multi-format document parsing.

Reference: FRD section 4.2 "Multi-Format Document Ingestion" and SRS
section 5.1 "Supported Input Formats".

Every parser here returns a `ParsedDocument` (plain text + structural
metadata) and nothing else - the caller hands `ParsedDocument.text` to
`app.services.ingestion.ingest_document`, so chunking, embedding, and
vector storage are entirely unaffected by which format came in.

Optional-dependency policy (same as app/services/vector_store.py): the
platform must stay runnable in offline/sandboxed environments, so no
parser is allowed to break import of this module. Instead:

  - .txt/.md  - stdlib only, always available.
  - .docx     - prefers `python-docx`, but falls back to a stdlib
                zipfile+ElementTree reader (a .docx *is* a zip containing
                word/document.xml), so DOCX ingestion works with zero
                external dependencies.
  - .pdf      - requires `pypdf`. There is no sane pure-python fallback
                for the PDF text layer, so a missing dependency raises
                `MissingParserDependency`, which the API surfaces as a
                503 rather than a 500.

Scanned/image-only PDFs and email (.eml/.msg) remain scaffolded - see
app/scaffold/parsers.py.
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from xml.etree import ElementTree

# WordprocessingML namespace used by .docx part word/document.xml.
_WORDPROCESSING_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# Collapse the runs of blank lines that PDF/DOCX extraction tends to emit,
# so chunking sees clean prose instead of whitespace padding.
_EXCESS_BLANK_LINES = re.compile(r"\n{3,}")


class ParserError(Exception):
    """Base class for recoverable, user-facing parsing failures."""


class UnsupportedFormat(ParserError):
    """The uploaded file extension has no parser wired up."""


class MissingParserDependency(ParserError):
    """A parser exists but its optional third-party package is not installed."""


class CorruptDocument(ParserError):
    """The file matched a supported format but could not be read."""


@dataclass
class ParsedDocument:
    """Plain text plus whatever structure the source format exposed.

    `metadata` is deliberately format-specific and free-form; it is carried
    for downstream consumers (document classification - see
    app/services/classification.py) and is not required by ingestion.
    """

    text: str
    source_type: str
    metadata: dict = field(default_factory=dict)


def _normalize(text: str) -> str:
    """Normalize line endings and collapse excessive blank runs."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _EXCESS_BLANK_LINES.sub("\n\n", text)
    return text.strip()


def parse_text(file_bytes: bytes) -> ParsedDocument:
    """Decode a plain-text/markdown upload.

    `errors="replace"` matches the pre-existing behaviour of
    `POST /api/v1/documents`: a stray non-UTF-8 byte degrades that one
    character rather than rejecting the whole document.
    """
    return ParsedDocument(
        text=_normalize(file_bytes.decode("utf-8", errors="replace")),
        source_type="text",
    )


def parse_pdf(file_bytes: bytes) -> ParsedDocument:
    """Extract the text layer of a PDF (FRD 4.2.1).

    Image-only/scanned pages have no text layer and come back empty; those
    need OCR, which is still scaffolded (FRD 4.2.2 - see
    app/scaffold/parsers.py). `metadata["empty_pages"]` records their
    1-based page numbers so the gap is visible rather than silent.
    """
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on install state
        raise MissingParserDependency(
            "PDF ingestion requires the `pypdf` package (pip install pypdf)."
        ) from exc

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:  # pypdf raises a wide range of parse errors
        raise CorruptDocument(f"Could not read PDF: {exc}") from exc

    empty_pages = [i for i, page_text in enumerate(pages, start=1) if not page_text.strip()]
    return ParsedDocument(
        text=_normalize("\n\n".join(pages)),
        source_type="pdf",
        metadata={"page_count": len(pages), "empty_pages": empty_pages},
    )


def parse_docx(file_bytes: bytes) -> ParsedDocument:
    """Extract text from a Word document (FRD 4.2.3).

    Uses `python-docx` when available so heading styles can be captured as
    metadata for downstream classification; otherwise falls back to reading
    word/document.xml directly, which yields the same paragraph text
    without the style information.
    """
    try:
        import docx  # python-docx
    except ImportError:
        return _parse_docx_stdlib(file_bytes)

    try:
        document = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in document.paragraphs]
        headings = [
            p.text.strip()
            for p in document.paragraphs
            if p.style is not None
            and p.style.name
            and p.style.name.startswith("Heading")
            and p.text.strip()
        ]
    except Exception as exc:
        raise CorruptDocument(f"Could not read DOCX: {exc}") from exc

    # Tables carry real content in most business documents; skipping them
    # would silently drop invoice/contract line items from the index.
    try:
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    paragraphs.append(" | ".join(cells))
    except Exception as exc:
        raise CorruptDocument(f"Could not read DOCX tables: {exc}") from exc

    return ParsedDocument(
        text=_normalize("\n".join(paragraphs)),
        source_type="docx",
        metadata={"headings": headings, "parser": "python-docx"},
    )


def _parse_docx_stdlib(file_bytes: bytes) -> ParsedDocument:
    """Zero-dependency .docx reader: a .docx is a zip holding XML parts.

    Each `<w:p>` is a paragraph and each `<w:t>` inside it is a text run, so
    joining runs per paragraph and paragraphs by newline reproduces the
    document's reading order for ordinary prose.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            xml_bytes = archive.read("word/document.xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise CorruptDocument(f"Could not read DOCX: {exc}") from exc

    try:
        root = ElementTree.fromstring(xml_bytes)
    except ElementTree.ParseError as exc:
        raise CorruptDocument(f"Could not parse DOCX XML: {exc}") from exc

    paragraphs: list[str] = []
    for paragraph in root.iter(f"{_WORDPROCESSING_NS}p"):
        runs = [node.text or "" for node in paragraph.iter(f"{_WORDPROCESSING_NS}t")]
        paragraphs.append("".join(runs))

    return ParsedDocument(
        text=_normalize("\n".join(paragraphs)),
        source_type="docx",
        metadata={"headings": [], "parser": "stdlib"},
    )


# Extension -> parser. Registering a new format here is the only change the
# API layer needs; `documents.py` reads this table for both dispatch and the
# "supported formats" error message.
PARSERS = {
    ".txt": parse_text,
    ".md": parse_text,
    ".pdf": parse_pdf,
    ".docx": parse_docx,
}

SUPPORTED_SUFFIXES: tuple[str, ...] = tuple(PARSERS)


def parse_upload(filename: str, file_bytes: bytes) -> ParsedDocument:
    """Dispatch an upload to the parser registered for its extension.

    Raises `UnsupportedFormat` for anything not in `PARSERS`; the legacy
    .doc binary format is called out separately because "rename it to .docx"
    is not a fix and users try it.
    """
    suffix = _suffix_of(filename)
    if suffix == ".doc":
        raise UnsupportedFormat(
            "Legacy .doc is not supported - re-save the file as .docx and retry."
        )

    parser = PARSERS.get(suffix)
    if parser is None:
        raise UnsupportedFormat(
            f"Unsupported file type '{suffix or filename}'. "
            f"Supported formats: {', '.join(SUPPORTED_SUFFIXES)}."
        )
    return parser(file_bytes)


def _suffix_of(filename: str) -> str:
    """Lowercased extension including the dot, or '' when there is none."""
    name = (filename or "").rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    if "." not in name.lstrip("."):
        return ""
    return "." + name.rsplit(".", 1)[-1].lower()
